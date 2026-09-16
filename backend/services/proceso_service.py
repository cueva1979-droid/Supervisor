import os
import re
import uuid
import pdfplumber
from datetime import datetime
from typing import Optional, Dict, List
from sqlalchemy.orm import Session
from models import ProcesoContratacion


def _fix_encoding(text: str) -> str:
    if not text:
        return text
    return text.replace('\ufffd', '').strip()


def extract_proceso_data(filepath: str) -> List[Dict]:
    full_text = ""
    with pdfplumber.open(filepath) as pdf:
        for page in pdf.pages:
            t = page.extract_text()
            if t:
                full_text += t + "\n"

    results = _extract_from_text(full_text)
    if results:
        return results
    return [_extract_single_process(full_text)]


def _extract_from_text(text: str) -> List[Dict]:
    """Extract multiple processes from text when PDF is a list/report."""
    results = []
    lines = text.split('\n')
    
    # Pattern for process codes like GADCCC-2026-001
    code_pattern = re.compile(r'(GADCCC-\d{4}-\d+)', re.IGNORECASE)
    
    # Find all lines with process codes
    for i, line in enumerate(lines):
        code_match = code_pattern.search(line)
        if not code_match:
            continue
        
        # Look back 1 line for the prefix (CPC-, FI-, LICS-, PE-, SIE-, VPN-)
        prefix = ""
        objeto_line = ""
        if i > 0:
            prev_line = lines[i - 1].strip()
            prefix_match = re.match(r'^([A-Z]{2,4})-', prev_line)
            if prefix_match:
                prefix = prefix_match.group(1)
                objeto_line = prev_line
        
        # Full code with prefix
        codigo = f"{prefix}-{code_match.group(1)}" if prefix else code_match.group(1)
        
        # Extract data from current and surrounding lines
        # Include lines: previous (prefix), current (code), next 2 lines
        block_lines = lines[max(0, i-1):min(len(lines), i+3)]
        block = '\n'.join(block_lines)
        
        # Also get the line with the code for objeto continuation (same line as GADCCC)
        objeto_extra = line.strip()
        
        proc = _parse_process_block(codigo, objeto_line, objeto_extra, block)
        if proc:
            results.append(proc)
    
    return results


def _parse_process_block(codigo: str, objeto_line: str, objeto_extra: str, block: str) -> Optional[Dict]:
    """Parse a text block that contains one process data."""
    
    # Skip non-process blocks
    if not codigo or 'Ingrese' in codigo:
        return None
    
    # Extract estado del proceso
    # Handle encoding issues: Ejecuci?n, Desierta, Adjudicada, etc.
    estado = ""
    estado_match = re.search(r'\b(Desiert[ao]|Adjudicad[ao]|Ejecuci\S*(?:\s*de\s+Contrato)?|Finalizad[ao]|Publicad[ao]|Suspendid[ao]|Recepci\S*(?:\s*de\s*)?|En\s+curso|Cancelad[ao]|Borrador)\b', block, re.IGNORECASE)
    if estado_match:
        raw = estado_match.group(1).strip()
        # Normalize before _fix_encoding removes replacement chars
        raw = re.sub(r'Ejecuci\S*n', 'Ejecución', raw)
        raw = re.sub(r'Ejecuci\S*n\s*de\s*Contrato', 'Ejecución de Contrato', raw)
        raw = re.sub(r'Recepci\S*n', 'Recepción', raw)
        estado = _fix_encoding(raw)
    
    # Extract presupuesto referencial
    presupuesto = None
    pres_match = re.search(r'\$[\s]*([\d,]+\.?\d*)', block)
    if pres_match:
        try:
            presupuesto = float(pres_match.group(1).replace(',', ''))
        except:
            pass
    
    # Extract fecha de publicación (format: 2026-03-20 / 17:30:00)
    fecha = ""
    fecha_match = re.search(r'(\d{4}-\d{2}-\d{2})\s*/\s*\d{2}:\d{2}', block)
    if fecha_match:
        fecha = fecha_match.group(1)
    else:
        fecha_match = re.search(r'(\d{4}-\d{2}-\d{2})', block)
        if fecha_match:
            fecha = fecha_match.group(1)
    
    # Extract objeto del proceso
    # Format:
    # Line N:   PREFIX- GOBIERNO AUTONOMO [OBJETO_PART1] Estado Provincia $...
    # Line N+1: GADCCC-YYYY-NNN DESCENTRALIZADO DEL [OBJETO_PART2] CHINCHIPE / HH:MM:SS
    # Line N+2: CANTON CENTINELA DEL ... CENTINELA  (entity part, NOT objeto)
    # Line N+3: CONDOR DEL CONDOR                    (entity part)
    
    objeto = ""
    
    # Get part1 from objeto_line (after "GOBIERNO AUTONOMO" and before estado)
    if objeto_line:
        clean = re.sub(r'^[A-Z]{2,4}-\s*', '', objeto_line).strip()
        obj_match = re.search(
            r'GOBIERNO AUTONOMO\s+(.*?)(?:\s+Desierta|\s+Adjudicada|\s+Ejecuci|\s+Finalizada|\s+Publicada|\s+Suspendida|\s+Recepci|\s+ZAMORA)',
            clean, re.IGNORECASE | re.DOTALL
        )
        if obj_match:
            objeto = _fix_encoding(obj_match.group(1))
    
    # Get part2 from objeto_extra (the GADCCC line, after "DESCENTRALIZADO DEL" and before CHINCHIPE/hora)
    if objeto_extra:
        clean_extra = re.sub(r'GADCCC-\d{4}-\d+\s*', '', objeto_extra).strip()
        clean_extra = re.sub(r'DESCENTRALIZADO DEL\s*', '', clean_extra).strip()
        # Remove everything from CHINCHIPE onwards
        clean_extra = re.split(r'\s+CHINCHIPE', clean_extra, flags=re.IGNORECASE)[0].strip()
        # Remove hora at the end
        clean_extra = re.sub(r'\s*/\s*\d{2}:\d{2}(:\d{2})?\s*$', '', clean_extra).strip()
        
        if clean_extra:
            objeto = (objeto + " " + clean_extra).strip()
    
    return {
        "codigo": codigo,
        "objeto_proceso": objeto,
        "estado_proceso": estado,
        "presupuesto_referencial": presupuesto,
        "fecha_publicacion": fecha,
    }


def _extract_single_process(text: str) -> Dict:
    """Fallback: extract single process data from text."""
    result = {
        "codigo": None,
        "objeto_proceso": None,
        "estado_proceso": None,
        "presupuesto_referencial": None,
        "fecha_publicacion": None,
    }
    
    patterns_code = [
        r'(?:C[ÓO]DIGO\s*(?:DEL)?\s*PROCESO)\s*[:#]?\s*([A-Z0-9\-/]{4,})',
        r'\b([A-Z]{2,5}-\d{4}-\d+)\b',
    ]
    for p in patterns_code:
        m = re.search(p, text, re.IGNORECASE)
        if m:
            result["codigo"] = _fix_encoding(m.group(1))
            break
    
    m = re.search(r'(?:OBJETO\s*(?:DEL)?\s*PROCESO)\s*[:#]?\s*(.+?)(?:\n|ESTADO|FECHA)', text, re.IGNORECASE)
    if m:
        result["objeto_proceso"] = _fix_encoding(m.group(1)[:500])
    
    for est in ['Desierta', 'Desierto', 'Adjudicada', 'Adjudicado', 'Ejecución de Contrato', 'Finalizada', 'Finalizado', 'Publicada', 'Publicado', 'Suspendida', 'Suspendido', 'Cancelada', 'Cancelado', 'Borrador']:
        if est.lower() in text.lower():
            result["estado_proceso"] = _fix_encoding(est)
            break
    
    m = re.search(r'(?:PRESUPUESTO\s*REFERENCIAL)\s*[:#]?\s*\$?\s*([\d.,]+)', text, re.IGNORECASE)
    if m:
        try:
            result["presupuesto_referencial"] = float(m.group(1).replace(',', '').replace('.', '', 1))
        except:
            pass
    
    m = re.search(r'(\d{4}-\d{2}-\d{2})', text)
    if m:
        result["fecha_publicacion"] = m.group(1)
    
    return result


def process_proceso_pdf(filepath: str, filename: str, db: Session) -> dict:
    all_procs = extract_proceso_data(filepath)

    if not all_procs or not any(p.get("codigo") for p in all_procs):
        raise ValueError("No se pudo extraer ningún código de proceso del documento")

    codigos_validos = [p["codigo"] for p in all_procs if p.get("codigo")]

    duplicados = (
        db.query(ProcesoContratacion.codigo)
        .filter(ProcesoContratacion.codigo.in_(codigos_validos))
        .all()
    )
    if duplicados:
        codigos_dup = [d.codigo for d in duplicados]
        raise ValueError(
            f"Los siguientes códigos ya existen en el sistema: {', '.join(codigos_dup)}. "
            "Elimine los procesos existentes antes de volver a cargar el archivo."
        )

    now = datetime.utcnow().isoformat()
    created_count = 0
    results = []

    for data in all_procs:
        if not data.get("codigo"):
            continue

        proc = ProcesoContratacion(
            id=str(uuid.uuid4()),
            filename=filename,
            codigo=data["codigo"],
            objeto_proceso=data.get("objeto_proceso"),
            estado_proceso=data.get("estado_proceso"),
            presupuesto_referencial=data.get("presupuesto_referencial"),
            fecha_publicacion=data.get("fecha_publicacion"),
            fecha_procesamiento=now,
        )
        db.add(proc)
        db.flush()
        db.refresh(proc)
        results.append(_proceso_to_dict(proc, created=True))
        created_count += 1

    return {
        "procesos_creados": created_count,
        "procesos_actualizados": 0,
        "total": created_count,
        "procesos": results,
        "mensaje": f"{created_count} proceso(s) creado(s)",
    }


def _proceso_to_dict(proc, created=False, updated=False) -> dict:
    result = {
        "id": proc.id,
        "filename": proc.filename or "",
        "codigo": proc.codigo,
        "objeto_proceso": proc.objeto_proceso or "",
        "estado_proceso": proc.estado_proceso or "",
        "presupuesto_referencial": proc.presupuesto_referencial,
        "fecha_publicacion": proc.fecha_publicacion or "",
        "fecha_procesamiento": proc.fecha_procesamiento or "",
    }
    if created:
        result["mensaje"] = "Proceso creado exitosamente"
    elif updated:
        result["mensaje"] = "Proceso actualizado exitosamente"
    return result


def list_procesos(db: Session) -> list:
    return [
        _proceso_to_dict(p)
        for p in db.query(ProcesoContratacion).order_by(ProcesoContratacion.fecha_procesamiento.desc()).all()
    ]


def get_proceso(proceso_id: str, db: Session) -> Optional[dict]:
    proc = db.query(ProcesoContratacion).filter(ProcesoContratacion.id == proceso_id).first()
    return _proceso_to_dict(proc) if proc else None


def delete_proceso(proceso_id: str, db: Session) -> bool:
    proc = db.query(ProcesoContratacion).filter(ProcesoContratacion.id == proceso_id).first()
    if not proc:
        return False
    db.delete(proc)
    db.commit()
    return True


def delete_all_procesos(db: Session) -> int:
    count = db.query(ProcesoContratacion).count()
    db.query(ProcesoContratacion).delete()
    db.commit()
    return count


def export_procesos_excel(db: Session) -> str:
    import os
    import re
    import tempfile
    import datetime
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from services.security import sanitize_excel

    procesos = db.query(ProcesoContratacion).order_by(
        ProcesoContratacion.fecha_procesamiento.desc()
    ).all()

    if not procesos:
        raise ValueError("No hay procesos de contratación para exportar")

    wb = Workbook()
    ws = wb.active
    ws.title = "Procesos"

    # Título
    ws.merge_cells("A1:G1")
    title_cell = ws.cell(row=1, column=1, value="Reporte de Procesos de Contratación")
    title_cell.font = Font(name="Calibri", bold=True, size=14, color="1F4E79")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    # Subtítulo
    now_str = datetime.datetime.now().strftime("%Y-%m-%d %H:%M")
    ws.merge_cells("A2:G2")
    sub_cell = ws.cell(row=2, column=1, value=f"Generado: {now_str}")
    sub_cell.font = Font(name="Calibri", size=10, italic=True, color="666666")
    sub_cell.alignment = Alignment(horizontal="center")

    # Headers
    headers = ["Código", "Objeto del Proceso", "Estado", "Presupuesto Referencial",
               "Fecha Publicación", "Archivo", "Fecha de Carga"]
    hfill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    hfont = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    halign = Alignment(horizontal="center", vertical="center")
    thin = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    for col, h in enumerate(headers, 1):
        c = ws.cell(row=3, column=col, value=h)
        c.fill = hfill
        c.font = hfont
        c.alignment = halign
        c.border = thin

    # Datos
    for i, proc in enumerate(procesos, 4):
        ws.cell(row=i, column=1, value=sanitize_excel(proc.codigo or ""))
        ws.cell(row=i, column=2, value=sanitize_excel(proc.objeto_proceso or ""))
        ws.cell(row=i, column=3, value=sanitize_excel(proc.estado_proceso or ""))
        ws.cell(row=i, column=4, value=proc.presupuesto_referencial)
        ws.cell(row=i, column=5, value=sanitize_excel(proc.fecha_publicacion or ""))
        ws.cell(row=i, column=6, value=sanitize_excel(proc.filename or ""))
        ws.cell(row=i, column=7, value=sanitize_excel(proc.fecha_procesamiento or ""))

        for col in range(1, 8):
            cell = ws.cell(row=i, column=col)
            cell.font = Font(name="Calibri", size=10)
            cell.border = thin
            if col == 2:
                cell.alignment = Alignment(wrap_text=True, vertical="center")
            elif col == 4:
                cell.number_format = '#,##0.00'
            else:
                cell.alignment = Alignment(vertical="center")

    # Auto-filter
    last_row = len(procesos) + 3
    ws.auto_filter.ref = f"A3:G{last_row}"

    # Anchos de columna
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 60
    ws.column_dimensions["C"].width = 20
    ws.column_dimensions["D"].width = 22
    ws.column_dimensions["E"].width = 16
    ws.column_dimensions["F"].width = 32
    ws.column_dimensions["G"].width = 22

    # Pie de página
    footer_row = last_row + 2
    ws.cell(row=footer_row, column=1,
            value=f"Reporte generado el: {now_str}").font = Font(
        name="Calibri", italic=True, size=9, color="666666")

    # Guardar
    fp = os.path.join(tempfile.gettempdir(), "Procesos_Contratacion.xlsx")
    wb.save(fp)
    return fp
